import sys, json
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE

path = sys.argv[1]
prs = Presentation(path)
print("slide size EMU:", prs.slide_width, prs.slide_height)
print("slide size in:", round(Emu(prs.slide_width).inches, 2), round(Emu(prs.slide_height).inches, 2))
print("slide count:", len(prs.slides))

# theme colors
try:
    theme = prs.slide_masters[0].part.part_related_by(
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme")
    xml = theme.blob.decode("utf-8", "ignore")
    import re
    colors = re.findall(r'<a:(dk1|lt1|dk2|lt2|accent1|accent2|accent3|accent4|accent5|accent6|hlink|folHlink)>.*?(?:val="([0-9A-Fa-f]{6})"|lastClr="([0-9A-Fa-f]{6})")', xml)
    print("\nTHEME COLORS:", colors[:16])
    fonts = re.findall(r'<a:(majorFont|minorFont)>.{0,400}?typeface="([^"]*)"', xml, re.S)
    print("THEME FONTS:", fonts)
except Exception as e:
    print("theme err", e)

for i, slide in enumerate(prs.slides, 1):
    print(f"\n===== Slide {i} =====")
    for sh in slide.shapes:
        try:
            pos = f"({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f}) {Emu(sh.width).inches:.2f}x{Emu(sh.height).inches:.2f}"
        except Exception:
            pos = "n/a"
        kind = str(sh.shape_type)
        print(f"  [{kind}] {sh.name} {pos}")
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            print("      IMG:", sh.image.filename or "", sh.image.size)
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                txt = "".join(r.text for r in p.runs)
                if txt.strip():
                    sizes = [r.font.size.pt if r.font.size else None for r in p.runs]
                    fonts = [r.font.name for r in p.runs]
                    cols = []
                    for r in p.runs:
                        try:
                            cols.append(str(r.font.color.rgb))
                        except Exception:
                            cols.append(None)
                    print(f"      T(lvl{p.level}) sz={sizes} font={fonts} color={cols}: {txt[:180]}")
        if sh.has_table:
            tb = sh.table
            print(f"      TABLE {len(tb.rows)}x{len(tb.columns)}")
            for r in tb.rows:
                print("        | " + " | ".join(c.text.replace("\n", " ")[:40] for c in r.cells))